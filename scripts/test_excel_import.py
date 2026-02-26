#!/usr/bin/env python
"""
Test de l'import Excel de véhicules (table risques).

Usage (depuis la racine du projet, avec l'environnement virtuel activé) :
    python scripts/test_excel_import.py

Le script :
  1. Génère un fichier Excel en mémoire avec 3 véhicules distincts + 1 doublon interne.
  2. Exécute l'import (même logique que la vue import_excel).
  3. Vérifie les compteurs créés / ignorés / erreurs.
  4. Réimporte le même fichier → tous ignorés.
  5. Nettoie les enregistrements de test.
"""
import os
import sys
import uuid
from io import BytesIO

# Force UTF-8 output (Windows console default is cp1252)
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# ── Setup Django ────────────────────────────────────────────────────────────
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
# Les apps sont enregistrées sans préfixe "apps." (settings.py ajoute apps/ au path)
sys.path.insert(0, os.path.join(ROOT, 'apps'))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')

import django
django.setup()

# ── Imports métier ───────────────────────────────────────────────────────────
import openpyxl
from django.utils import timezone
from contrats.models import Risques
from contrats.views import _safe_str, _safe_int

SEP = "-" * 62

# ── Immatriculations de test (préfixe TEST_ pour faciliter le nettoyage) ────
TEST_IMMATS = ['TEST-N01-AA', 'TEST-N02-BB', 'TEST-N03-CC']


# ────────────────────────────────────────────────────────────────────────────

def create_test_excel() -> BytesIO:
    """Crée un fichier Excel en mémoire avec 4 lignes (dont 1 doublon interne)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Véhicules"

    headers = [
        'veh_immat', 'veh_marque', 'veh_modele', 'veh_chassis',
        'veh_type', 'veh_puissance', 'veh_nbplace',
        'veh_cat', 'veh_energie', 'veh_usage',
    ]
    ws.append(headers)

    rows = [
        # immat          marque      modele    chassis      type      pf   pl   cat   energie   usage
        ['TEST-N01-AA', 'Toyota',   'Corolla', 'CH-AA-001', 'Berline', 7,  5,  'VL', 'Essence',          'Personnel'],
        ['TEST-N02-BB', 'Mercedes', 'Sprinter','CH-BB-002', 'Fourgon', 12, 15, 'TC', 'Electrique',           'Transport en commun'],
        ['TEST-N01-AA', 'Toyota',   'Corolla', 'CH-AA-001', 'Berline', 7,  5,  'VL', 'Essence',          'Personnel'],  # doublon interne
        ['TEST-N03-CC', 'Honda',    'CR-V',    'CH-CC-003', 'SUV',     9,  7,  'VL', 'Hybride',          'Commercial'],
    ]
    for r in rows:
        ws.append(r)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def run_import(excel_bytes: BytesIO) -> dict:
    """
    Reproduit exactement la logique de RisquesViewSet.import_excel()
    sans passer par HTTP.
    """
    wb = openpyxl.load_workbook(excel_bytes, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    if len(rows) < 2:
        return {'created': [], 'skipped': [], 'errors': []}

    headers = [
        str(h).strip().lower().replace(' ', '_') if h is not None else ''
        for h in rows[0]
    ]

    created, skipped, errors = [], [], []

    for row_values in rows[1:]:
        row = dict(zip(headers, row_values))

        immat = _safe_str(row.get('veh_immat'))
        if not immat:
            continue
        immat = immat.upper()

        if Risques.objects.filter(veh_immat__iexact=immat, effacer=False).exists():
            skipped.append(immat)
            continue

        try:
            Risques.objects.create(
                id_risque=str(uuid.uuid4()),
                type_risque='Véhicule',
                veh_immat=immat,
                veh_marque=_safe_str(row.get('veh_marque')),
                veh_modele=_safe_str(row.get('veh_modele')),
                veh_chassis=_safe_str(row.get('veh_chassis')),
                veh_type=_safe_str(row.get('veh_type')),
                veh_puissance=_safe_str(row.get('veh_puissance')),
                veh_nbplace=_safe_int(row.get('veh_nbplace')),
                veh_cat=_safe_str(row.get('veh_cat')),
                veh_energie=_safe_str(row.get('veh_energie')),
                veh_usage=_safe_str(row.get('veh_usage')),
                veh_valeur_neuve=_safe_str(
                    row.get('veh_valeur_neuve') or row.get('veh_valeurneuve')
                ),
                veh_valeur_venale=_safe_str(
                    row.get('veh_valeur_venale') or row.get('veh_valeurvenale')
                ),
                veh_nbremorque=_safe_int(row.get('veh_nbremorque')),
                effacer=False,
                sync=False,
                date_enreg=timezone.now(),
            )
            created.append(immat)
        except Exception as e:
            errors.append({'immat': immat, 'error': str(e)})

    return {'created': created, 'skipped': skipped, 'errors': errors}


def print_result(result: dict):
    print(f"  Créés   ({len(result['created'])}) : {result['created']}")
    print(f"  Ignorés ({len(result['skipped'])}) : {result['skipped']}")
    if result['errors']:
        print(f"  Erreurs ({len(result['errors'])}) : {result['errors']}")


def cleanup():
    n, _ = Risques.objects.filter(veh_immat__in=TEST_IMMATS).delete()
    return n


def assert_eq(label, actual, expected):
    if actual != expected:
        raise AssertionError(f"[FAIL] {label} : attendu {expected}, obtenu {actual}")
    print(f"  ✓ {label} = {actual}")


# ────────────────────────────────────────────────────────────────────────────

def main():
    print(SEP)
    print("  TEST : Import Excel — table risques")
    print(SEP)

    # Précaution : nettoyer les éventuels restes d'un test précédent
    n_pre = cleanup()
    if n_pre:
        print(f"\n  (nettoyage préalable : {n_pre} enregistrement(s) supprimé(s))")

    # ── Passe 1 : premier import ─────────────────────────────────────────────
    print("\n[1/3] Premier import (3 véhicules distincts + 1 doublon dans le fichier)…")
    r1 = run_import(create_test_excel())
    print_result(r1)
    assert_eq("créés",   len(r1['created']), 3)
    assert_eq("ignorés", len(r1['skipped']), 1)   # doublon interne (TEST-N01-AA vu 2x)
    assert_eq("erreurs", len(r1['errors']),  0)

    # ── Passe 2 : réimport → tous doublons ──────────────────────────────────
    print("\n[2/3] Réimport du même fichier (tous déjà en base → tous ignorés)…")
    r2 = run_import(create_test_excel())
    print_result(r2)
    assert_eq("créés",   len(r2['created']), 0)
    assert_eq("ignorés", len(r2['skipped']), 4)   # 4 lignes (dont le doublon interne) déjà en base
    assert_eq("erreurs", len(r2['errors']),  0)

    # ── Passe 3 : nettoyage ──────────────────────────────────────────────────
    print("\n[3/3] Nettoyage des enregistrements de test…")
    n = cleanup()
    print(f"  {n} enregistrement(s) supprimé(s).")
    assert_eq("supprimés", n, 3)

    print(f"\n{SEP}")
    print("  Tous les tests sont passés ✓")
    print(SEP)


if __name__ == '__main__':
    main()
